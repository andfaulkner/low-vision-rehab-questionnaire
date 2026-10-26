#!/usr/bin/env python3
"""Draw the PRISMA 2020 flow diagram (Page et al. 2021) from the screening
workbook's Counts tab, a JSON file, or hand overrides, and write the numbers
as a Markdown table with ready-to-paste Results sentences.

Inputs (pick one; --override applies on top):
  --counts-csv data/LVR_screening_sheet_2_Counts.csv   (default)
  --xlsx data/LVR_screening_sheet_2.xlsx               (read-only, cached values)
  --json counts.json                                   (see --json-schema)

The script checks the arithmetic of the diagram (identified - removed =
screened; screened - excluded = sought; sought - not retrieved = assessed;
assessed - excluded = included) and prints every discrepancy. It still draws
the diagram, but stamps it DRAFT and exits with status 2, so a diagram whose
numbers do not add up is never pasted into the paper unnoticed.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import textwrap
from pathlib import Path

# --------------------------------------------------------------------------
# Canonical counts dictionary
# --------------------------------------------------------------------------

JSON_SCHEMA = {
    "_comment": "Example input for prisma_flow.py --json. Keys you leave out are treated as missing.",
    "databases": {"PubMed": 78, "Scopus": 88},
    "registers": {},
    "other_sources": {"Key sources from the study outline": 8},
    "duplicates_removed": 39,
    "automation_removed": 0,
    "other_removed": 0,
    "records_screened": 135,
    "screened_so_far": 135,
    "records_excluded": 110,
    "ta_exclusion_reasons": {"No questionnaire/screening item": 60, "Not about low vision/visual function": 50},
    "reports_sought": 25,
    "reports_not_retrieved": 1,
    "reports_assessed": 24,
    "reports_excluded": {"No accuracy/validation data": 10, "Other (see notes)": 2},
    "reports_included": 12,
    "studies_included": 11,
    "other": {
        "_comment": "Only used with --other-mode separate (official right-hand column).",
        "sought": 8, "not_retrieved": 0, "assessed": 8,
        "excluded": {"Review/opinion/editorial": 2}, "included": 6,
    },
}

INT_KEYS = [
    "duplicates_removed", "automation_removed", "other_removed",
    "records_screened", "screened_so_far", "records_excluded",
    "reports_sought", "reports_not_retrieved", "reports_assessed",
    "reports_included", "studies_included",
]
DICT_KEYS = ["databases", "registers", "other_sources",
             "ta_exclusion_reasons", "reports_excluded"]


def empty_counts() -> dict:
    c = {k: None for k in INT_KEYS}
    c.update({k: {} for k in DICT_KEYS})
    c["automation_removed"] = 0
    c["other_removed"] = 0
    c["other"] = {}
    c["stated_total_identified"] = None
    c["notes"] = []
    return c


# --------------------------------------------------------------------------
# Readers
# --------------------------------------------------------------------------

def _norm(label: str) -> str:
    s = label.strip().replace("–", "-").replace("—", "-").replace("→", "->")
    return re.sub(r"\s+", " ", s).lower()


def _to_int(value) -> int | None:
    if value is None:
        return None
    s = str(value).strip()
    if s == "":
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def rows_from_csv(path: Path) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return [row for row in csv.reader(fh)]


def rows_from_xlsx(path: Path, sheet: str) -> list[list[str]]:
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is not installed: python3 -m pip install --user "
                 "--break-system-packages openpyxl  (or use --counts-csv)")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f"sheet {sheet!r} not in {path}; sheets: {wb.sheetnames}")
    rows = []
    for row in wb[sheet].iter_rows(values_only=True):
        rows.append(["" if v is None else str(v) for v in row])
    return rows


def parse_counts_rows(rows: list[list[str]], not_retrieved_reason: str) -> dict:
    """Parse the label / value / note layout of the Counts tab."""
    c = empty_counts()
    section = None
    for row in rows:
        row = list(row) + ["", "", ""]
        label, v1, v2 = row[0], row[1], row[2]
        key = _norm(label)
        if key == "":
            if section == "reasons":
                section = None
            continue
        if key.startswith("exclusion reasons"):
            section = "reasons-header"
            continue
        if section == "reasons-header" and key.startswith("reason"):
            section = "reasons"
            continue
        if section == "reasons":
            ta, ft = _to_int(v1), _to_int(v2)
            reason = label.strip()
            if ta is not None:
                c["ta_exclusion_reasons"][reason] = ta
            if ft is not None:
                c["reports_excluded"][reason] = ft
            continue
        n = _to_int(v1)
        m = re.match(r"^records identified - (.+)$", key)
        if m:
            if n is None:
                continue
            name = re.split(r"\s[-–—]\s", label.strip(), maxsplit=1)[-1].strip()
            name = name[:1].upper() + name[1:]
            if "other" in m.group(1):
                c["other_sources"][name] = n
            elif "register" in m.group(1):
                c["registers"][name] = n
            else:
                c["databases"][name] = n
        elif key.startswith("total records identified"):
            c["stated_total_identified"] = n
        elif key.startswith("duplicates removed"):
            c["duplicates_removed"] = n
        elif key.startswith("records screened"):
            c["records_screened"] = n
        elif key.startswith("screened so far"):
            c["screened_so_far"] = n
        elif key.startswith("excluded at title/abstract"):
            c["records_excluded"] = n
        elif key.startswith("included + maybe") or key.endswith("-> full text"):
            c["reports_sought"] = n
        elif key.startswith("full texts assessed"):
            c["reports_assessed"] = n
        elif key.startswith("excluded at full text"):
            c["_ft_excluded_total"] = n
        elif key.startswith("studies included"):
            c["studies_included"] = n
    # PRISMA 2020 counts reports that could not be obtained as "not retrieved",
    # not as "excluded"; the workbook records them as an exclusion reason.
    nr_key = next((k for k in c["reports_excluded"] if _norm(k) == _norm(not_retrieved_reason)), None)
    if nr_key is not None:
        nr = c["reports_excluded"].pop(nr_key)
        c["reports_not_retrieved"] = nr
        if nr:
            if c["reports_assessed"] is not None:
                c["reports_assessed"] -= nr
            c["notes"].append(
                f"{nr} report(s) with reason '{nr_key}' were moved from 'excluded' to 'not retrieved' "
                "and subtracted from 'assessed for eligibility' (PRISMA 2020 box definitions).")
    else:
        c["reports_not_retrieved"] = 0
    if c["reports_included"] is None:
        c["reports_included"] = c["studies_included"]
        c["notes"].append("The workbook counts included reports; studies = reports unless you override "
                          "studies_included (several reports of one study, or one report of several studies).")
    c["reports_excluded"] = {k: v for k, v in c["reports_excluded"].items() if v}
    c["ta_exclusion_reasons"] = {k: v for k, v in c["ta_exclusion_reasons"].items() if v}
    return c


def counts_from_json(path: Path) -> dict:
    c = empty_counts()
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        sys.exit(f"{path}: expected a JSON object (see --json-schema)")
    for k, v in data.items():
        if k.startswith("_"):
            continue
        if k in DICT_KEYS:
            if isinstance(v, int):
                v = {"Other sources": v}
            c[k] = {str(a): int(b) for a, b in v.items() if not str(a).startswith("_")}
        elif k in INT_KEYS or k == "stated_total_identified":
            c[k] = _to_int(v)
        elif k == "other":
            c["other"] = {a: b for a, b in v.items() if not a.startswith("_")}
        else:
            c["notes"].append(f"ignored unknown JSON key {k!r}")
    if c["reports_not_retrieved"] is None:
        c["reports_not_retrieved"] = 0
    if c["reports_included"] is None and c["studies_included"] is not None:
        c["reports_included"] = c["studies_included"]
    if c["studies_included"] is None and c["reports_included"] is not None:
        c["studies_included"] = c["reports_included"]
    return c


def apply_overrides(c: dict, overrides: list[str]) -> None:
    for item in overrides:
        if "=" not in item:
            sys.exit(f"--override expects key=value, got {item!r}")
        key, _, value = item.partition("=")
        key = key.strip()
        n = _to_int(value)
        if n is None:
            sys.exit(f"--override {key}: value {value!r} is not an integer")
        if key.startswith("other.excluded."):
            c["other"].setdefault("excluded", {})[key[len("other.excluded."):]] = n
        elif key.startswith("other."):
            c["other"][key[len("other."):]] = n
        elif "." in key:
            top, _, sub = key.partition(".")
            if top not in DICT_KEYS:
                sys.exit(f"--override: unknown dictionary {top!r}; known: {', '.join(DICT_KEYS)}")
            if n == 0:
                c[top].pop(sub, None)
            else:
                c[top][sub] = n
        elif key in INT_KEYS or key == "stated_total_identified":
            c[key] = n
        else:
            sys.exit(f"--override: unknown key {key!r}; known: {', '.join(INT_KEYS + DICT_KEYS)}")


# --------------------------------------------------------------------------
# Arithmetic checks
# --------------------------------------------------------------------------

def check_counts(c: dict, other_mode: str) -> tuple[list[str], list[str]]:
    """Return (discrepancies, warnings). Discrepancies are arithmetic failures
    or missing boxes; warnings are things the reviewer should know."""
    disc, warn = [], []
    g = c.get
    db = sum(c["databases"].values()) + sum(c["registers"].values())
    other = sum(c["other_sources"].values())
    identified_left = db + (other if other_mode == "same-pipeline" else 0)

    if not c["databases"] and not c["registers"]:
        disc.append("no database or register counts found")
    if c["stated_total_identified"] is not None and c["stated_total_identified"] != db + other:
        disc.append(f"stated total identified ({c['stated_total_identified']}) != databases + registers + "
                    f"other sources ({db + other})")

    removed = (g("duplicates_removed") or 0) + (g("automation_removed") or 0) + (g("other_removed") or 0)
    if g("records_screened") is None:
        disc.append("records_screened is missing")
    elif g("duplicates_removed") is None:
        disc.append("duplicates_removed is missing")
    elif identified_left - removed != g("records_screened"):
        disc.append(f"identified ({identified_left}) - removed before screening ({removed}) = "
                    f"{identified_left - removed}, but records screened = {g('records_screened')}")

    partial = (g("screened_so_far") is not None and g("records_screened") is not None
               and g("screened_so_far") < g("records_screened"))
    if partial:
        warn.append(f"screening is incomplete: {g('screened_so_far')} of {g('records_screened')} records have "
                    "a title/abstract decision, so everything below 'Records screened' is provisional")

    if g("records_excluded") is None or g("reports_sought") is None:
        disc.append("records_excluded or reports_sought is missing")
    elif g("records_screened") is not None and \
            g("records_screened") - g("records_excluded") != g("reports_sought"):
        disc.append(f"records screened ({g('records_screened')}) - excluded ({g('records_excluded')}) = "
                    f"{g('records_screened') - g('records_excluded')}, but reports sought = {g('reports_sought')}")

    ta_sum = sum(c["ta_exclusion_reasons"].values())
    if c["ta_exclusion_reasons"] and g("records_excluded") is not None and ta_sum != g("records_excluded"):
        disc.append(f"title/abstract exclusion reasons sum to {ta_sum}, but records excluded = {g('records_excluded')}")

    if g("reports_assessed") is None:
        disc.append("reports_assessed is missing")
    elif g("reports_sought") is not None and \
            g("reports_sought") - (g("reports_not_retrieved") or 0) != g("reports_assessed"):
        disc.append(f"reports sought ({g('reports_sought')}) - not retrieved ({g('reports_not_retrieved') or 0}) = "
                    f"{g('reports_sought') - (g('reports_not_retrieved') or 0)}, but reports assessed = "
                    f"{g('reports_assessed')}")

    ft_sum = sum(c["reports_excluded"].values())
    if c.get("_ft_excluded_total") is not None and \
            ft_sum + (g("reports_not_retrieved") or 0) != c["_ft_excluded_total"]:
        disc.append(f"full-text exclusion reasons sum to {ft_sum} (+ {g('reports_not_retrieved') or 0} not retrieved), "
                    f"but the workbook's 'Excluded at full text' = {c['_ft_excluded_total']}")
    if g("reports_included") is None:
        disc.append("reports_included / studies_included is missing")
    elif g("reports_assessed") is not None:
        inc_left = g("reports_assessed") - ft_sum
        other_inc = (c["other"].get("included") or 0) if other_mode == "separate" else 0
        if inc_left + other_inc != g("reports_included"):
            disc.append(f"reports assessed ({g('reports_assessed')}) - excluded ({ft_sum})"
                        + (f" + other-method reports included ({other_inc})" if other_mode == "separate" else "")
                        + f" = {inc_left + other_inc}, but reports included = {g('reports_included')}")

    if other_mode == "separate" and other:
        o = c["other"]
        for k in ("sought", "not_retrieved", "assessed", "included"):
            if o.get(k) is None:
                disc.append(f"--other-mode separate needs other.{k} (use --override other.{k}=N or the JSON 'other' block)")
        if all(o.get(k) is not None for k in ("sought", "not_retrieved", "assessed")):
            if other != o["sought"]:
                warn.append(f"other methods: identified ({other}) != sought for retrieval ({o['sought']}); "
                            "PRISMA 2020 expects these to match unless some records were dropped before retrieval")
            if o["sought"] - o["not_retrieved"] != o["assessed"]:
                disc.append(f"other methods: sought ({o['sought']}) - not retrieved ({o['not_retrieved']}) != "
                            f"assessed ({o['assessed']})")
            if o.get("included") is not None:
                o_ex = sum((o.get("excluded") or {}).values())
                if o["assessed"] - o_ex != o["included"]:
                    disc.append(f"other methods: assessed ({o['assessed']}) - excluded ({o_ex}) != included ({o['included']})")
    return disc, warn


# --------------------------------------------------------------------------
# Drawing
# --------------------------------------------------------------------------

def _n(v) -> str:
    return "?" if v is None else str(v)


def _lines(text: str, width: int) -> list[str]:
    """Wrap each paragraph; '(n = X)' is glued with non-breaking spaces so it never splits."""
    out = []
    text = re.sub(r"\(n = ([^)]*)\)", lambda m: "(n = " + m.group(1).replace(" ", " ") + ")", text)
    for para in text.split("\n"):
        if para == "":
            out.append("")
            continue
        indent = len(para) - len(para.lstrip(" "))
        out.extend(textwrap.wrap(para, width=width, subsequent_indent=" " * (indent + 2)) or [""])
    return out


class Diagram:
    """Boxes and arrows on a plain inch-based canvas (matplotlib, Agg)."""

    LINE_H = 0.165      # inches per text line at 8.5 pt
    PAD = 0.12
    CHARS_PER_IN = 15.5

    def __init__(self, width: float, height: float):
        try:
            import matplotlib
        except ImportError:  # python3 -I hides the user site; re-add that one trusted dir
            import site
            sys.path.append(site.getusersitepackages())
            try:
                import matplotlib
            except ImportError:
                raise SystemExit("ERROR: matplotlib could not be imported. "
                                 "install with: python3 -m pip install --user --break-system-packages matplotlib")
        matplotlib.use("Agg")
        matplotlib.rcParams["svg.hashsalt"] = "prisma-flow"
        import matplotlib.pyplot as plt
        from matplotlib.patches import FancyArrowPatch, Rectangle
        self.plt, self.Rect, self.Arrow = plt, Rectangle, FancyArrowPatch
        self.fig, self.ax = plt.subplots(figsize=(width, height))
        self.ax.set_xlim(0, width)
        self.ax.set_ylim(0, height)
        self.ax.set_aspect("equal")
        self.ax.axis("off")
        self.fig.subplots_adjust(0, 0, 1, 1)
        self.W, self.H = width, height

    def box(self, x: float, top: float, w: float, text: str, *, shade: str = "white",
            min_h: float = 0.45, align: str = "left") -> dict:
        lines = _lines(text, int(w * self.CHARS_PER_IN))
        h = max(min_h, len(lines) * self.LINE_H + 2 * self.PAD)
        self.ax.add_patch(self.Rect((x, top - h), w, h, fc=shade, ec="black", lw=0.9))
        tx = x + self.PAD if align == "left" else x + w / 2
        self.ax.text(tx, top - h / 2, "\n".join(lines), fontsize=8.5, va="center",
                     ha=align, family="DejaVu Sans", linespacing=1.25)
        return {"x": x, "y": top - h, "w": w, "h": h, "top": top, "cx": x + w / 2, "cy": top - h / 2}

    def arrow(self, p0, p1, *, via=None):
        pts = [p0] + (via or []) + [p1]
        for a, b in zip(pts[:-2], pts[1:-1]):
            self.ax.plot([a[0], b[0]], [a[1], b[1]], color="black", lw=0.9, solid_capstyle="round")
        self.ax.add_patch(self.Arrow(pts[-2], pts[-1], arrowstyle="-|>", mutation_scale=9,
                                     color="black", lw=0.9, shrinkA=0, shrinkB=0))

    def down(self, b_from: dict, b_to: dict):
        self.arrow((b_from["cx"], b_from["y"]), (b_to["cx"], b_to["top"]))

    def right(self, b_from: dict, b_to: dict):
        self.arrow((b_from["x"] + b_from["w"], b_from["cy"]), (b_to["x"], b_from["cy"]))

    def side_label(self, x: float, y0: float, y1: float, text: str):
        self.ax.add_patch(self.Rect((x, y0), 0.42, y1 - y0, fc="#dbe7f3", ec="black", lw=0.9))
        self.ax.text(x + 0.21, (y0 + y1) / 2, text, rotation=90, fontsize=9, va="center",
                     ha="center", family="DejaVu Sans")

    def heading(self, x: float, top: float, w: float, text: str) -> dict:
        return self.box(x, top, w, text, shade="#dbe7f3", align="center", min_h=0.45)

    def banner(self, text: str):
        self.ax.text(self.W / 2, self.H - 0.1, text, color="#b00020", fontsize=10, ha="center",
                     va="top", weight="bold", family="DejaVu Sans")

    def footnote(self, text: str):
        lines = textwrap.wrap(text, width=int((self.W - 0.6) * self.CHARS_PER_IN * 1.1))
        self.ax.text(0.3, self.ax.get_ylim()[0] + 0.1, "\n".join(lines), fontsize=7.5, ha="left",
                     va="bottom", family="DejaVu Sans")

    def crop(self, bottom: float):
        """Shrink the canvas so that it ends just below the lowest element."""
        self.fig.set_size_inches(self.W, self.H - bottom)
        self.ax.set_ylim(bottom, self.H)

    def save(self, svg: Path, png: Path, dpi: int):
        self.fig.savefig(svg, format="svg", metadata={"Date": None, "Creator": None})
        self.fig.savefig(png, format="png", dpi=dpi, metadata={"Software": None})
        self.plt.close(self.fig)


def _sources_text(prefix: str, d: dict, fallback: str) -> str:
    if not d:
        return f"{prefix}\n  {fallback} (n = 0)"
    return prefix + "".join(f"\n  {k} (n = {v})" for k, v in d.items())


def _reasons_text(prefix: str, d: dict) -> str:
    if not d:
        return f"{prefix} (n = 0)"
    return prefix + ":" + "".join(f"\n  {k} (n = {v})" for k, v in d.items())


def draw(c: dict, other_mode: str, out_svg: Path, out_png: Path, *, banner: str | None,
         dpi: int, footnote: str | None) -> None:
    two_col = sum(c["other_sources"].values()) > 0
    col_w, gap, left_x = 2.9, 0.3, 0.9
    excl_x = left_x + col_w + gap
    other_x = excl_x + col_w + 0.55
    other_excl_x = other_x + col_w + gap
    other_cols = 2 if other_mode == "separate" else 1   # same-pipeline has only the identification box
    if two_col:
        width = other_x + other_cols * col_w + (other_cols - 1) * gap + 0.3
    else:
        width = excl_x + col_w + 0.3
    height = 12.0   # tall scratch canvas; cropped to the content at the end
    d = Diagram(width, height)
    top = height - (0.45 if banner else 0.25)

    h_left = d.heading(left_x, top, 2 * col_w + gap, "Identification of studies via databases and registers")
    if two_col:
        d.heading(other_x, top, other_cols * col_w + (other_cols - 1) * gap,
                  "Identification of studies via other methods")
    y = h_left["y"] - 0.3

    # Identification row
    ident_txt = _sources_text("Records identified from:", c["databases"], "Databases")
    ident_txt += "".join(f"\n  {k} (n = {v})" for k, v in c["registers"].items())
    b_ident = d.box(left_x, y, col_w, ident_txt)
    removed_txt = ("Records removed before screening:"
                   f"\n  Duplicate records removed (n = {_n(c['duplicates_removed'])})"
                   f"\n  Records marked as ineligible by automation tools (n = {_n(c['automation_removed'])})"
                   f"\n  Records removed for other reasons (n = {_n(c['other_removed'])})")
    b_removed = d.box(excl_x, y, col_w, removed_txt)
    d.right(b_ident, b_removed)
    b_other = None
    if two_col:
        b_other = d.box(other_x, y, col_w,
                        _sources_text("Records identified from:", c["other_sources"], "Other sources"))
    ident_bottom = min(b_ident["y"], b_removed["y"], b_other["y"] if b_other else y)
    y = ident_bottom - 0.45

    # Screening rows
    b_screen = d.box(left_x, y, col_w, f"Records screened\n(n = {_n(c['records_screened'])})")
    b_excl = d.box(excl_x, y, col_w, f"Records excluded\n(n = {_n(c['records_excluded'])})")
    d.down(b_ident, b_screen)
    d.right(b_screen, b_excl)
    if two_col and other_mode == "same-pipeline":
        # Hand-added records went through the same title/abstract screening.
        mid_y = ident_bottom - 0.22   # below the whole identification row, so it clears the 'removed' box
        d.arrow((b_other["cx"], b_other["y"]), (b_screen["cx"] + 0.7, b_screen["top"]),
                via=[(b_other["cx"], mid_y), (b_screen["cx"] + 0.7, mid_y)])
    y = b_screen["y"] - 0.4
    b_sought = d.box(left_x, y, col_w, f"Reports sought for retrieval\n(n = {_n(c['reports_sought'])})")
    b_nr = d.box(excl_x, y, col_w, f"Reports not retrieved\n(n = {_n(c['reports_not_retrieved'])})")
    d.down(b_screen, b_sought)
    d.right(b_sought, b_nr)
    b_osought = None
    if two_col and other_mode == "separate":
        o = c["other"]
        b_osought = d.box(other_x, y, col_w, f"Reports sought for retrieval\n(n = {_n(o.get('sought'))})")
        b_onr = d.box(other_excl_x, y, col_w, f"Reports not retrieved\n(n = {_n(o.get('not_retrieved'))})")
        d.down(b_other, b_osought)
        d.right(b_osought, b_onr)
    y = b_sought["y"] - 0.4
    b_assess = d.box(left_x, y, col_w, f"Reports assessed for eligibility\n(n = {_n(c['reports_assessed'])})")
    b_ftex = d.box(excl_x, y, col_w, _reasons_text("Reports excluded", c["reports_excluded"]))
    d.down(b_sought, b_assess)
    d.right(b_assess, b_ftex)
    b_oassess = None
    if b_osought is not None:
        o = c["other"]
        b_oassess = d.box(other_x, y, col_w, f"Reports assessed for eligibility\n(n = {_n(o.get('assessed'))})")
        b_oex = d.box(other_excl_x, y, col_w, _reasons_text("Reports excluded", o.get("excluded") or {}))
        d.down(b_osought, b_oassess)
        d.right(b_oassess, b_oex)
    screen_bottom = min(b_assess["y"], b_ftex["y"], b_oassess["y"] if b_oassess else y)
    y = screen_bottom - 0.45

    # Included
    b_inc = d.box(left_x, y, col_w,
                  f"Studies included in review\n(n = {_n(c['studies_included'])})\n"
                  f"Reports of included studies\n(n = {_n(c['reports_included'])})")
    d.down(b_assess, b_inc)
    if b_oassess is not None:
        mid_y = (b_oassess["y"] + b_inc["top"]) / 2
        d.arrow((b_oassess["cx"], b_oassess["y"]), (b_inc["x"] + b_inc["w"], b_inc["cy"]),
                via=[(b_oassess["cx"], mid_y), (b_inc["x"] + b_inc["w"] + 0.5, mid_y),
                     (b_inc["x"] + b_inc["w"] + 0.5, b_inc["cy"])])

    d.side_label(0.3, ident_bottom, h_left["y"] - 0.1, "Identification")
    d.side_label(0.3, screen_bottom, b_screen["top"] + 0.1, "Screening")
    d.side_label(0.3, b_inc["y"], b_inc["top"] + 0.1, "Included")

    bottom = b_inc["y"] - 0.3
    if footnote:
        foot_lines = len(textwrap.wrap(footnote, width=int((width - 0.6) * d.CHARS_PER_IN * 1.1)))
        bottom -= foot_lines * 0.14 + 0.1
    d.crop(bottom)
    if banner:
        d.banner(banner)
    if footnote:
        d.footnote(footnote)
    d.save(out_svg, out_png, dpi)


# --------------------------------------------------------------------------
# Markdown output
# --------------------------------------------------------------------------

def _words(n: int | None) -> str:
    small = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten"]
    if n is None:
        return "{n}"
    return small[n] if 0 <= n <= 10 else str(n)


def _plural(n, one: str, many: str) -> str:
    return one if n == 1 else many


def counts_markdown(c: dict, other_mode: str, disc: list[str], warn: list[str], source: str,
                    figure_path: str) -> str:
    db_total = sum(c["databases"].values()) + sum(c["registers"].values())
    other_total = sum(c["other_sources"].values())
    total = db_total + other_total
    ft_sum = sum(c["reports_excluded"].values())
    rows = [("Records identified from databases and registers", db_total,
             "; ".join(f"{k} {v}" for k, v in {**c["databases"], **c["registers"]}.items()))]
    if other_total:
        rows.append(("Records identified from other methods", other_total,
                     "; ".join(f"{k} {v}" for k, v in c["other_sources"].items())))
    rows += [
        ("Total records identified", total, ""),
        ("Duplicate records removed", c["duplicates_removed"], ""),
        ("Records removed by automation tools", c["automation_removed"], ""),
        ("Records removed for other reasons", c["other_removed"], ""),
        ("Records screened (title/abstract)", c["records_screened"], ""),
        ("Records excluded at title/abstract", c["records_excluded"], ""),
        ("Reports sought for retrieval", c["reports_sought"], ""),
        ("Reports not retrieved", c["reports_not_retrieved"], ""),
        ("Reports assessed for eligibility", c["reports_assessed"], ""),
        ("Reports excluded at full text", ft_sum, ""),
        ("Reports of included studies", c["reports_included"], ""),
        ("Studies included in review", c["studies_included"], ""),
    ]
    if other_mode == "separate" and other_total:
        o = c["other"]
        rows += [("Other methods: reports sought", o.get("sought"), ""),
                 ("Other methods: reports not retrieved", o.get("not_retrieved"), ""),
                 ("Other methods: reports assessed", o.get("assessed"), ""),
                 ("Other methods: reports excluded", sum((o.get("excluded") or {}).values()), ""),
                 ("Other methods: reports included", o.get("included"), "")]

    md = ["# PRISMA 2020 flow diagram counts", "",
          f"Source: `{source}`. Diagram: `{figure_path}`. Other-methods mode: `{other_mode}`.", ""]
    if disc:
        md += ["> **DRAFT: the counts do not reconcile.** Fix the workbook or pass `--override` "
               "before using these numbers.", ""]
        md += [f"> - {x}" for x in disc] + [""]
    if warn:
        md += ["> Notes:", ""] + [f"> - {x}" for x in warn] + [""]
    md += ["| Stage | n | Detail |", "|---|---:|---|"]
    md += [f"| {a} | {_n(b)} | {cc} |" for a, b, cc in rows]
    md += [""]
    if c["ta_exclusion_reasons"] or c["reports_excluded"]:
        md += ["## Exclusion reasons", "", "| Reason | Title/abstract | Full text |", "|---|---:|---:|"]
        for r in sorted(set(c["ta_exclusion_reasons"]) | set(c["reports_excluded"])):
            md.append(f"| {r} | {c['ta_exclusion_reasons'].get(r, 0)} | {c['reports_excluded'].get(r, 0)} |")
        md += [""]
    for note in c.get("notes", []):
        md.append(f"- {note}")
    if c.get("notes"):
        md.append("")

    # Results sentences
    src = "; ".join(f"{k}, n = {v}" for k, v in {**c["databases"], **c["registers"]}.items())
    s = [f"The searches identified {db_total} records ({src})"]
    if other_total:
        s[0] += (f", and {other_total} further {_plural(other_total, 'record was', 'records were')} identified "
                 f"from other sources ({', '.join(c['other_sources'])})")
    s[0] += "."
    removed = c["duplicates_removed"]
    s.append(f"After {_n(removed)} duplicate {_plural(removed, 'record was', 'records were')} removed, "
             f"{_n(c['records_screened'])} records were screened by title and abstract and "
             f"{_n(c['records_excluded'])} were excluded.")
    nr = c["reports_not_retrieved"] or 0
    sought = c["reports_sought"]
    s.append(f"The full {_plural(sought, 'text', 'texts')} of {_n(sought)} {_plural(sought, 'report was', 'reports were')} sought"
             + (f"; {nr} could not be retrieved" if nr else "")
             + f", and {_n(c['reports_assessed'])} {_plural(c['reports_assessed'], 'was', 'were')} assessed for eligibility.")
    if ft_sum:
        reasons = "; ".join(f"{k[0].lower() + k[1:]}, n = {v}" for k, v in c["reports_excluded"].items())
        s.append(f"Of these, {ft_sum} {_plural(ft_sum, 'was', 'were')} excluded ({reasons}).")
    st, rp = c["studies_included"], c["reports_included"]
    s.append(f"In total, {_words(st)} {_plural(st, 'study', 'studies')} ({_n(rp)} {_plural(rp, 'report', 'reports')}) "
             f"{_plural(st, 'was', 'were')} included in the review (Figure {{N}}).")
    md += ["## Results text", "", " ".join(s), ""]
    if any("incomplete" in w for w in warn):
        md += ["_Screening is still in progress; regenerate this file when the workbook is complete._", ""]
    return "\n".join(md)


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = p.add_mutually_exclusive_group()
    src.add_argument("--counts-csv", type=Path, default=Path("data/LVR_screening_sheet_2_Counts.csv"),
                     help="Counts tab exported as CSV (default: %(default)s)")
    src.add_argument("--xlsx", type=Path, help="read the Counts sheet from the workbook (read-only, cached values)")
    src.add_argument("--json", type=Path, help="counts as JSON (see --json-schema)")
    p.add_argument("--sheet", default="Counts", help="sheet name for --xlsx (default: %(default)s)")
    p.add_argument("--override", action="append", default=[], metavar="KEY=N",
                   help="fix one number by hand, e.g. reports_not_retrieved=2, databases.PubMed=80, "
                        "'reports_excluded.Other (see notes)=1', other.sought=8; repeatable; 0 removes a dictionary entry")
    p.add_argument("--other-mode", choices=["same-pipeline", "separate"], default="same-pipeline",
                   help="how records from other sources flow: 'same-pipeline' (they were title/abstract-screened "
                        "with the database records, as the workbook does) or 'separate' (official PRISMA 2020 "
                        "right-hand column; needs other.sought/not_retrieved/assessed/included)")
    p.add_argument("--not-retrieved-reason", default="Full text not available",
                   help="full-text exclusion reason that PRISMA counts as 'reports not retrieved' (default: %(default)s)")
    p.add_argument("--out-dir", type=Path, default=Path("paper/figures"),
                   help="where the SVG and PNG go (default: %(default)s)")
    p.add_argument("--basename", default="prisma_flow", help="file stem for the diagram (default: %(default)s)")
    p.add_argument("--counts-md", type=Path, default=Path("paper/prisma_counts.md"),
                   help="Markdown table + Results sentences (default: %(default)s)")
    p.add_argument("--footnote", default=None, help="footnote text printed under the diagram")
    p.add_argument("--dpi", type=int, default=200)
    p.add_argument("--no-draft-banner", action="store_true",
                   help="do not stamp DRAFT on a diagram whose counts do not reconcile")
    p.add_argument("--json-schema", action="store_true", help="print an example JSON input and exit")
    p.add_argument("--dump-json", type=Path, help="also write the parsed counts (after overrides) to this JSON file")
    a = p.parse_args(argv)

    if a.json_schema:
        print(json.dumps(JSON_SCHEMA, indent=2, ensure_ascii=False))
        return 0

    if a.json:
        if not a.json.exists():
            sys.exit(f"{a.json} not found")
        c, source = counts_from_json(a.json), str(a.json)
    elif a.xlsx:
        if not a.xlsx.exists():
            sys.exit(f"{a.xlsx} not found")
        c, source = parse_counts_rows(rows_from_xlsx(a.xlsx, a.sheet), a.not_retrieved_reason), f"{a.xlsx}#{a.sheet}"
    else:
        if not a.counts_csv.exists():
            sys.exit(f"{a.counts_csv} not found (regenerate the CSV export or pass --xlsx / --json)")
        c, source = parse_counts_rows(rows_from_csv(a.counts_csv), a.not_retrieved_reason), str(a.counts_csv)
    apply_overrides(c, a.override)

    if a.other_mode == "same-pipeline" and sum(c["other_sources"].values()) > 0:
        c["notes"].append("Records identified from other sources were title/abstract-screened together with the "
                          "database records (--other-mode same-pipeline), so the arrow from the other-methods box "
                          "joins 'Records screened'. Say so in the figure legend.")

    disc, warn = check_counts(c, a.other_mode)
    for w in warn:
        print(f"note: {w}", file=sys.stderr)
    for x in disc:
        print(f"DISCREPANCY: {x}", file=sys.stderr)

    a.out_dir.mkdir(parents=True, exist_ok=True)
    a.counts_md.parent.mkdir(parents=True, exist_ok=True)
    svg, png = a.out_dir / f"{a.basename}.svg", a.out_dir / f"{a.basename}.png"
    banner = None
    if disc and not a.no_draft_banner:
        banner = f"DRAFT: {len(disc)} count(s) do not reconcile; see {a.counts_md.name}"
    draw(c, a.other_mode, svg, png, banner=banner, dpi=a.dpi, footnote=a.footnote)
    a.counts_md.write_text(counts_markdown(c, a.other_mode, disc, warn, source, str(png)), encoding="utf-8")
    if a.dump_json:
        dump = {k: v for k, v in c.items() if not k.startswith("_") and k != "notes"}
        a.dump_json.write_text(json.dumps(dump, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {svg}, {png}, {a.counts_md}" + (f" ({len(disc)} discrepancies, exit 2)" if disc else ""))
    return 2 if disc else 0


if __name__ == "__main__":
    sys.exit(main())
